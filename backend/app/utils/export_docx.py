from docx import Document
from docx.shared import Inches
from typing import List


def _get_title(section):
    if isinstance(section, dict):
        return section.get("title", "Untitled Section")
    return getattr(section, "title", "Untitled Section")


def _get_content(section):
    if isinstance(section, dict):
        return section.get("content", "")
    return getattr(section, "content", "")


def export_to_docx(project_title: str, sections: List) -> bytes:
    doc = Document()

    # Add title
    title = doc.add_heading(project_title, 0)
    title.alignment = 1

    # Add sections
    for section in sections:
        sec_title = _get_title(section)
        sec_content = _get_content(section)

        doc.add_heading(sec_title, level=1)

        if sec_content:
            lines = sec_content.split('\n')
            for line in lines:
                line = line.strip()
                if line:
                    if line.startswith('-') or line.startswith('•'):
                        p = doc.add_paragraph(style='List Bullet')
                        p.add_run(line[1:].strip())
                    else:
                        doc.add_paragraph(line)

        doc.add_paragraph()

    import io
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
