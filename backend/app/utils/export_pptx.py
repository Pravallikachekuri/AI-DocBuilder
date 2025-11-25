from pptx import Presentation
from pptx.util import Pt
from typing import List


def _get_title(section):
    if isinstance(section, dict):
        return section.get("title", "Untitled Section")
    return getattr(section, "title", "Untitled Section")


def _get_content(section):
    if isinstance(section, dict):
        return section.get("content", "")
    return getattr(section, "content", "")


def export_to_pptx(project_title: str, sections: List) -> bytes:
    prs = Presentation()

    # Title slide
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = project_title
    slide.placeholders[1].text = "Generated with AI Doc Platform"

    # Content slides
    for section in sections:
        sec_title = _get_title(section)
        sec_content = _get_content(section)

        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)

        slide.shapes.title.text = sec_title

        tf = slide.placeholders[1].text_frame
        tf.clear()

        if sec_content:
            lines = sec_content.split('\n')

            for i, line in enumerate(lines):
                line = line.strip()
                if not line:
                    continue

                if line.startswith('-') or line.startswith('•'):
                    line = line[1:].strip()

                if i == 0:
                    p = tf.paragraphs[0]
                else:
                    p = tf.add_paragraph()

                p.text = line
                p.font.size = Pt(18)

    import io
    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
